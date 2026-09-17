from pathlib import Path
from decimal import Decimal

from openpyxl import load_workbook

import pytest

import workers.rfb_cnpj.exporter as exporter
from workers.rfb_cnpj.exporter import write_csv, write_xlsx
from workers.rfb_cnpj.models import CnpjFilters, CnpjRecord


def test_write_csv_uses_utf8_bom_and_blocks_sensitive_fields(tmp_path: Path):
    output = tmp_path / "base.csv"
    result = write_csv([CnpjRecord(cnpj="1", razao_social="=Empresa Teste")], ("cnpj", "razao_social"), output)

    assert result.row_count == 1
    content = output.read_bytes()
    assert content.startswith(b"\xef\xbb\xbf")
    assert b";" in content
    assert b"\r\n" in content
    assert "'=Empresa Teste" in output.read_text(encoding="utf-8-sig")


def test_write_xlsx_creates_required_operational_sheets(tmp_path: Path):
    output = tmp_path / "base.xlsx"
    result = write_xlsx(
        [CnpjRecord(cnpj="1", razao_social="+Empresa Teste", cnae_principal="4321500", municipio="Campinas", uf="SP", capital_social=Decimal("1234.56"))],
        ("cnpj", "razao_social", "municipio", "uf", "capital_social"),
        output,
        CnpjFilters(segment="Energia solar", uf="SP"),
    )

    workbook = load_workbook(output)
    assert result.format == "xlsx"
    assert workbook.sheetnames == ["Leads", "Resumo", "Filtros aplicados", "Leia-me"]
    assert workbook["Leads"]["A1"].value == "cnpj"
    assert workbook["Leads"]["B2"].value == "'+Empresa Teste"
    assert workbook["Leads"]["E2"].value == 1234.56
    assert workbook["Resumo"]["A2"].value == "protocolo"
    assert workbook["Filtros aplicados"]["A2"].value == "segmento"
    workbook.close()


@pytest.mark.parametrize("dangerous", [" =SUM(1,1)", "\t@cmd", "\r-HYPERLINK('x')", "+1+1"])
def test_csv_neutralizes_formula_injection_after_whitespace(tmp_path: Path, dangerous: str):
    output = tmp_path / "formula.csv"

    write_csv([CnpjRecord(cnpj="00000000000191", razao_social=dangerous)], ("cnpj", "razao_social"), output)

    content = output.read_text(encoding="utf-8-sig")
    assert "00000000000191" in content
    assert ";'" in content


def test_csv_failure_removes_temporary_and_final_files(tmp_path: Path, monkeypatch):
    output = tmp_path / "base.csv"

    def fail_validation(*_args, **_kwargs):
        raise RuntimeError("falha simulada")

    monkeypatch.setattr(exporter, "_validate_csv", fail_validation)
    with pytest.raises(RuntimeError, match="falha simulada"):
        write_csv([CnpjRecord(cnpj="00000000000191", razao_social="Empresa")], ("cnpj", "razao_social"), output)

    assert not output.exists()
    assert list(tmp_path.iterdir()) == []
